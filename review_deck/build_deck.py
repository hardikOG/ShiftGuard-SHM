"""Fill the Zeroth-Review template structure with the Review-2 ShiftGuard content.

  python build_deck.py      (reads structure.pptx, writes ShiftGuard_Review2.pptx)
"""
import copy

from pptx import Presentation
from pptx.util import Inches
from lxml import etree

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
DATE = "October 2026"
TWIN_IMG = "../docs/screenshots/4_damage_cropped.png"
TIMELINE_IMG = "../results/figures/timeline.png"

prs = Presentation("structure.pptx")
S = list(prs.slides)


def shape(slide, name_prefix):
    return next(sh for sh in slide.shapes if sh.name.startswith(name_prefix))


def set_text(sh, items):
    """items: list of (text, template_paragraph_index). Each new paragraph is a copy of the
    template paragraph at that index (keeps bullet, font, size, colour); only its text changes."""
    tb = sh.text_frame._txBody
    olds = list(tb.iter(A + "p"))
    for _, idx in items:
        assert idx < len(olds), (sh.name, idx, len(olds))
    protos = [copy.deepcopy(p) for p in olds]
    for p in olds:
        tb.remove(p)
    for text, idx in items:
        p = copy.deepcopy(protos[idx])
        runs = p.findall(A + "r")
        for extra in p.findall(A + "br") + p.findall(A + "fld"):
            p.remove(extra)
        if not runs:                              # paragraph had no run: build one
            r = etree.SubElement(p, A + "r"); etree.SubElement(r, A + "rPr", lang="en-US")
            etree.SubElement(r, A + "t"); runs = [r]
        for extra in runs[1:]:
            p.remove(extra)
        runs[0].find(A + "t").text = text
        end = p.find(A + "endParaRPr")
        if end is not None:                       # endParaRPr must stay last
            p.remove(end); p.append(end)
        tb.append(p)


def footer(slide, n):
    for sh in slide.shapes:
        if sh.name.startswith("Date Placeholder"):
            set_text(sh, [(DATE, 0)])
        elif sh.name.startswith("Slide Number Placeholder"):
            set_text(sh, [(str(n), 0)])


def roomy(sh, size_pt, space_pt=10, align=None):
    """Bigger text and more space between bullets, for slides with few, short points."""
    for p in sh.text_frame._txBody.iter(A + "p"):
        ppr = p.find(A + "pPr")
        if ppr is None:
            ppr = etree.SubElement(p, A + "pPr"); p.insert(0, ppr)
        if align:
            ppr.set("algn", align)
        for old in ppr.findall(A + "spcBef"):
            ppr.remove(old)
        spc = etree.Element(A + "spcBef"); etree.SubElement(spc, A + "spcPts", val=str(space_pt * 100))
        ppr.insert(0, spc)
        heading = any(r.find(A + "rPr") is not None and r.find(A + "rPr").get("b") == "1" for r in p.findall(A + "r"))
        for r in p.findall(A + "r"):
            r.find(A + "rPr").set("sz", str(int((size_pt + (1 if heading else 0)) * 100)))


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def fill_like(dst, src):
    """Give shape dst the fill colour of shape src."""
    sp_d, sp_s = dst._element.spPr, src._element.spPr
    for f in sp_d.findall(A + "solidFill"):
        sp_d.remove(f)
    f = copy.deepcopy(sp_s.find(A + "solidFill"))
    geom = sp_d.find(A + "prstGeom")
    geom.addnext(f)


def solid(dst, hexcolor):
    sp = dst._element.spPr
    for f in sp.findall(A + "solidFill"):
        sp.remove(f)
    f = etree.Element(A + "solidFill"); etree.SubElement(f, A + "srgbClr", val=hexcolor)
    sp.find(A + "prstGeom").addnext(f)


B, H = 0, 3        # paragraph indices on the Problem-slide body: bullet, blue subheading

# ---------------------------------------------------------------- 1 title
s = S[0]
set_text(shape(s, "Title 1"), [("ShiftGuard: AIoT Structural Health Monitoring", 0)])
sub = shape(s, "Subtitle 2")
n_sub = len(sub.text_frame.paragraphs)
texts = [p.text for p in sub.text_frame.paragraphs]
new = []
for i, t in enumerate(texts):
    if t.startswith("2. Team"): t = "2. Advait Anand (RA Number)"
    elif t.startswith("3. Team"): t = "3. Akash Kumar (RA Number)"
    elif t.startswith("4. Team"): t = "4. Yograj (RA Number)"
    elif "Zeroth Review" in t: t = f"Review 2 | {DATE}"
    new.append((t, i))
set_text(sub, new)
notes(s, "We are building a system that watches a bridge 24/7 with IoT sensors and AI, and when something "
         "looks wrong it tells you WHAT is wrong: real damage, unusual weather, or just a broken sensor.")

# ---------------------------------------------------------------- 2 problem
s = S[1]; footer(s, 2)
set_text(shape(s, "Title 1"), [("The Problem We Found", 0)])
set_text(shape(s, "Content"), [
    ("In the Zeroth Review, our model judged every sensor reading on its own, allowing a 5 % error rate per reading.", B),
    ("At one reading per hour, that is 5 % of 8,760 hours ≈ 440 false alarms a year. Engineers would soon ignore it.", B),
    ("Every alarm was called “damage”. But a bridge's readings also change when it is very cold "
     "(Z24's deck stiffens when it freezes) or when an IoT sensor breaks (drift, noise, dropout, glitches).", B),
    ("A real deployment cannot tell us which sensor broke or when, so this is hard to test on real hardware alone.", B),
    ("So our question became", H),
    ("When the alarm rings: is it damage, weather, or a broken sensor? And can we promise false alarms are rare?", B),
])
roomy(shape(s, "Content"), 22, 12)
notes(s, "Key point: a 5 % error looks small, but over a whole year of hourly checks it becomes hundreds of false "
         "alarms. And the old system could not say WHY it raised an alarm. Cold weather and broken sensors look "
         "like damage to a simple model.")

# ---------------------------------------------------------------- 3 idea
s = S[2]; footer(s, 3)
set_text(shape(s, "Title 1"), [("Our Solution: ShiftGuard", 0)])
set_text(shape(s, "Content"), [
    ("Learn only what a healthy bridge looks like: predict each sensor from the other sensors and the temperature.", B),
    ("Hide the most suspicious sensor and check again. A broken sensor disagrees on its own; real damage makes a whole region disagree.", B),
    ("Collect evidence over time instead of judging one reading. One odd hour does nothing; a real change keeps adding up.", B),
    ("Three detectors run side by side (structure, each sensor, weather), so every alarm names its cause.", B),
    ("The guarantee", H),
    ("On a healthy bridge, the expected time to a false alarm is at least 10,000 hours (about 14 months).", B),
])
roomy(shape(s, "Content"), 22, 12)
notes(s, "Think of it as a team of sensors cross-checking each other. If one sensor says something the others "
         "don't agree with, that sensor is the problem. If a whole area of the bridge changes together, the "
         "structure changed. And instead of panicking at one strange reading, we wait until the evidence adds up.")

# ---------------------------------------------------------------- 4 workflow
s = S[3]; footer(s, 4)
set_text(shape(s, "TitleBox"), [("Work Flow of ShiftGuard", 0)])
box = {sh.name: sh for sh in s.shapes}
flow = {
    "FlowBox101": ["Virtual IoT sensors: 8 accelerometers + 1 thermometer, over MQTT every hour"],
    "FlowBox103": ["Feature extraction: 16 frequency bands per sensor"],
    "FlowBox105": ["Masked-sensor model", "predicts each sensor from the others + temperature"],
    "FlowBox106": ["Mask & impute", "hides the most suspicious sensor, fills lost packets"],
    "FlowBox108": ["Conformal p-values: compared only with healthy hours at the same temperature"],
    "FlowBox110": ["Three e-detectors collect evidence: Structure · Each sensor · Weather"],
    "FlowBox112": ["Alarm only when evidence ≥ 10,000 → expected time to a false alarm ≥ 10,000 h"],
    "FlowBox114": ["Structural damage", "with severity: Moderate / Severe"],
    "FlowBox115": ["Broken sensor", "names which sensor"],
    "FlowBox116": ["Unusual weather", "logged, not called damage"],
    "FlowBox118": ["Health score + damage location"],
    "FlowBox120": ["Dashboard + live digital twin"],
    "FlowBox122": ["Alert → maintenance engineer", "email via Node-RED for damage and broken sensors"],
    "Caption123": ["Trained only on healthy data. Every alarm says what changed: the structure, a sensor, or the weather."],
}
for name, lines in flow.items():
    set_text(box[name], [(t, i) for i, t in enumerate(lines)])
solid(box["FlowBox114"], "C0392B")    # damage = red
solid(box["FlowBox115"], "D9822B")    # broken sensor = orange
solid(box["FlowBox116"], "2A78D6")    # weather = blue
notes(s, "Read it top to bottom: sensors → features → the cross-checking model → fair comparison with healthy "
         "hours at the same temperature → evidence collectors → one of three verdicts → dashboard and email.")

# ---------------------------------------------------------------- 5 how it works
s = S[4]; footer(s, 5)
set_text(shape(s, "Title 1"), [("How It Works (In Plain Words)", 0)])
set_text(shape(s, "Content"), [
    ("Listen: 8 virtual accelerometers and a thermometer send a 20-second vibration recording every hour.", 0),
    ("Fingerprint: each recording becomes 16 frequency-band values per sensor.", 0),
    ("Cross-check: the model guesses each sensor from the others; the gap is that sensor's “disagreement score”.", 0),
    ("Compare fairly: today's score is ranked only against healthy hours at a similar temperature (conformal p-value).", 0),
    ("Add up evidence: small surprises accumulate; normal hours slowly erase them (CUSUM-type e-detector).", 0),
    ("Decide and explain: damage, broken sensor or weather, plus severity, shown on the dashboard and emailed.", 0),
])
roomy(shape(s, "Content"), 22, 14, align="l")
notes(s, "If asked 'what is conformal prediction?': it turns any score into a probability-like p-value by ranking "
         "it against scores from healthy data, with no assumptions about the distribution. 'What is an e-detector?': "
         "it multiplies evidence from each hour; under no change its expected value stays at 1, which is what gives "
         "the false-alarm guarantee.")

# ---------------------------------------------------------------- 6 testbed (+ screenshot)
s = S[5]; footer(s, 6)
set_text(shape(s, "Title 1"), [("Software-Defined AIoT Testbed", 0)])
body = shape(s, "Content")
body.width = Inches(6.3)
set_text(body, [
    ("Virtual sensor network over MQTT: a Tinkercad/Wokwi-style circuit, but for a whole bridge.", B),
    ("A fault injector breaks sensors on demand, so we know exactly which sensor failed and when.", B),
    ("Edge monitor: 4.9 ms per reading on a laptop.", B),
    ("Live digital twin: click a sensor to break it, or add snow, damage, or a repair.", B),
    ("Self-healing: restarts crashed parts and fills lost packets.", B),
    ("Optional: Wokwi ESP32 + DHT22 as a real temperature node.", B),
])
roomy(body, 20, 10, align="l")
pic = s.shapes.add_picture(TWIN_IMG, Inches(6.75), Inches(1.35), width=Inches(6.2))
pic.name = "Twin screenshot"
cap = s.shapes.add_textbox(Inches(6.75), Inches(1.35) + pic.height + Inches(0.05), Inches(6.2), Inches(0.35))
cap.name = "Twin caption"
cap.text_frame.text = "Live digital twin during severe damage: located at S5–S6, severity Severe (live demo)"
r = cap.text_frame.paragraphs[0].runs[0]
r.font.size = Inches(0.155); r.font.italic = True; r.font.name = "Times New Roman"
notes(s, "Demo: run 'python run_demo.py'. Break sensor S3 → it says 'broken sensor S3', not damage. "
         "Cold snap → 'unusual weather'. Damage → 'structural damage near S5 and S6'.")

# ---------------------------------------------------------------- 7 results (3 columns)
s = S[6]; footer(s, 7)
set_text(shape(s, "Title 1"), [("Results", 0)])
cols = [
    (shape(s, "Content"), "Simulated year", [
        "False alarms: 26 h vs 347 h (old method)",
        "Damage found in under 4 hours",
        "Broken sensor blamed on the right sensor: 94 %",
        "Faults or cold snap called damage: 0 %"]),
    (shape(s, "Column50"), "Real Z24 bridge data", [
        "False alarms: 10 h vs 391 h (old method)",
        "Faults called damage: 0.3 % (old: 23 %)",
        "Cold snap called damage: 0 % (old: 16.5 %)",
        "Severe damage detected; moderate missed"]),
    (shape(s, "Column51"), "What each part adds", [
        "Without evidence collection: ≈ 330 false-alarm episodes",
        "Without sensor masking: 91 % of faults called damage",
        "Without the weather detector: 33 % of the cold snap called damage",
        "Without temperature matching: 4× more false-alarm hours"]),
]
for sh, head, bullets in cols:
    set_text(sh, [(head, 0)] + [(b, 1) for b in bullets])
    roomy(sh, 18, 12, align="l")
set_text(shape(s, "FinalStatement"), [("ShiftGuard cuts false alarms by more than 10× and tells damage, weather and "
                                       "broken sensors apart, on simulated and real bridge data.", 0)])
notes(s, "Simulation numbers are the MLP model, averaged over 3 runs; Z24 numbers are the ridge model, 3 runs. "
         "Be upfront: on real Z24 data, moderate damage is not yet caught; that is our next step.")

# ---------------------------------------------------------------- 8 timeline figure
s = S[7]; footer(s, 8)
set_text(shape(s, "Title 1"), [("One Monitored Year at a Glance", 0)])
body = shape(s, "Content")
body.top, body.height = Inches(6.25), Inches(0.6)
set_text(body, [("Top: weather and injected events. Middle: the old method fires all year. "
                 "Bottom: ShiftGuard raises few alarms, and each names its cause.", B)])
img_h = Inches(5.35)
pic = s.shapes.add_picture(TIMELINE_IMG, 0, Inches(0.85), height=img_h)
pic.left = int((prs.slide_width - pic.width) / 2)
pic.name = "Timeline figure"
notes(s, "Blue = weather, orange = broken sensor, red = damage. The grey ticks in the second row are the old "
         "method's alarms: they happen all year. ShiftGuard's lines only cross the dashed threshold during real events.")

# ---------------------------------------------------------------- 9 limitations & next steps
s = S[8]; footer(s, 9)
set_text(shape(s, "Title 1"), [("Limitations and Next Steps", 0)])
set_text(shape(s, "Content"), [
    ("Honest limitations", H),
    ("On real Z24 data, only severe damage is detected so far; moderate damage is missed.", B),
    ("The guarantee assumes healthy hours are interchangeable. Real hours are correlated, so we see a few false-alarm episodes a year, not zero.", B),
    ("Handles one broken sensor at a time.", B),
    ("Next steps", H),
    ("Use the full Z24 sensor layouts and modal (frequency) features to catch moderate damage.", B),
    ("Put a real ESP32 sensor node in the loop.", B),
    ("Finalise and submit the paper (IEEE format).", B),
])
roomy(shape(s, "Content"), 21, 9, align="l")
notes(s, "Saying the limitations ourselves makes the work more credible. The moderate-damage gap is a sensor-"
         "placement / feature issue on Z24's shared channels: the old method is no better there.")

footer(S[9], 10)
prs.save("ShiftGuard_Review2.pptx")
print("saved ShiftGuard_Review2.pptx")
